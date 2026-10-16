"""Regenerate the DOCX and PDF sample résumés in input/sample_resumes/."""

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLES_DIR = PROJECT_ROOT / "input" / "sample_resumes"

EMMA_PDF_LINES = [
    "Emma Novak",
    "Python Developer",
    "emma.novak@example.com | Vienna, Austria",
    "",
    "SUMMARY",
    "Python developer with 5+ years of professional experience building",
    "backend services and web applications.",
    "",
    "SKILLS",
    "Python, Django, REST APIs, SQL, PostgreSQL, Docker, Redis, Git",
    "",
    "EXPERIENCE",
    "Python Developer - DataWorks (2021 - present)",
    "- Developed REST APIs with Django for a logistics platform",
    "- Deployed services with Docker and managed PostgreSQL schemas",
    "- Wrote complex SQL reports and data migrations",
    "",
    "Backend Developer - SoftLine (2019 - 2021)",
    "- Built Python scripts for data processing and automation",
    "- Maintained internal web applications used by 200+ staff",
    "",
    "EDUCATION",
    "B.Sc. Software Engineering, TU Wien",
]

DAVID_DOCX_PARAGRAPHS = [
    ("David Okafor", "Title"),
    ("Python Developer", "Normal"),
    ("david.okafor@example.com | Lagos, Nigeria", "Normal"),
    ("SUMMARY", "Heading 2"),
    ("Software developer with 5 years of experience delivering web "
     "applications for the healthcare industry.", "Normal"),
    ("SKILLS", "Heading 2"),
    ("Python, Django, SQL Server, MySQL, Docker, REST APIs, Git", "Normal"),
    ("EXPERIENCE", "Heading 2"),
    ("Python Developer - HealthSoft (2020 - present)", "Normal"),
    ("Built Django applications exposing REST APIs to mobile clients", "List Bullet"),
    ("Containerized services with Docker for consistent deployments", "List Bullet"),
    ("Designed SQL Server schemas and optimized slow queries", "List Bullet"),
    ("Junior Developer - WebStart (2018 - 2020)", "Normal"),
    ("Maintained PHP and Python internal tools", "List Bullet"),
    ("EDUCATION", "Heading 2"),
    ("B.Sc. Computer Science, University of Lagos", "Normal"),
]


def build_pdf(lines):
    """Assemble a minimal single-page PDF with plain text lines."""
    content = ["BT", "/F1 10 Tf", "40 750 Td", "14 TL"]
    for line in lines:
        escaped = (
            line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        )
        content.append(f"({escaped}) Tj T*")
    content.append("ET")
    stream = "\n".join(content).encode("latin-1")

    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
        4: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        5: (b"<< /Length " + str(len(stream)).encode()
            + b" >>\nstream\n" + stream + b"\nendstream"),
    }

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objects):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + objects[num] + b"\nendobj\n"

    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for num in sorted(objects):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF").encode()
    return bytes(out)


def make_pdf_sample() -> Path:
    target = SAMPLES_DIR / "emma_pdf_candidate.pdf"
    target.write_bytes(build_pdf(EMMA_PDF_LINES))
    return target


def make_docx_sample() -> Path:
    from docx import Document

    doc = Document()
    for text, style in DAVID_DOCX_PARAGRAPHS:
        if style == "Title":
            doc.add_heading(text, level=0)
        else:
            doc.add_paragraph(text, style=style)
    target = SAMPLES_DIR / "david_docx_candidate.docx"
    doc.save(target)
    return target


if __name__ == "__main__":
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    created = [make_pdf_sample(), make_docx_sample()]
    for path in created:
        print(f"created {path.relative_to(PROJECT_ROOT)}")
    sys.exit(0)
